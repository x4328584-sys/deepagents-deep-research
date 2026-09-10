"""Safe, bounded extraction for supported uploaded document formats."""

from __future__ import annotations

import asyncio
import json
from pathlib import Path
from typing import Any, cast
from zipfile import BadZipFile, ZipFile

import pandas as pd
from docx import Document
from langchain_core.tools import BaseTool, StructuredTool
from openpyxl import load_workbook
from pypdf import PdfReader

from api.context import get_session_dir
from utils.config import Settings
from utils.path_utils import safe_existing_file
from utils.security import ALLOWED_UPLOAD_EXTENSIONS, SecurityError


def _file_error(code: str, message: str) -> dict[str, Any]:
    return {"ok": False, "error": {"code": code, "message": message, "retryable": False}}


def _truncate(content: str, limit: int) -> tuple[str, bool]:
    if len(content) <= limit:
        return content, False
    return f"{content[:limit]}\n\n[Content truncated at configured extraction limit]", True


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_json(path: Path) -> str:
    data = json.loads(_read_text(path))
    return json.dumps(data, ensure_ascii=False, indent=2)


def _read_jsonl(path: Path) -> str:
    records = []
    for line_number, line in enumerate(_read_text(path).splitlines(), start=1):
        if not line.strip():
            continue
        try:
            records.append(json.dumps(json.loads(line), ensure_ascii=False))
        except json.JSONDecodeError as exc:
            raise ValueError(f"invalid JSONL at line {line_number}") from exc
    return "\n".join(records)


def _read_csv(path: Path) -> str:
    frame = pd.read_csv(path, nrows=5_000)
    return cast(str, frame.to_csv(index=False))


def _read_docx(path: Path) -> str:
    document = Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_number, table in enumerate(document.tables, start=1):
        blocks.append(f"[Table {table_number}]")
        blocks.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(blocks)


def _read_pdf(path: Path) -> str:
    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("encrypted PDF cannot be read") from exc
    pages = []
    for number, page in enumerate(reader.pages[:500], start=1):
        pages.append(f"[Page {number}]\n{page.extract_text() or ''}")
    return "\n\n".join(pages)


def _read_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for sheet in workbook.worksheets:
            sections.append(f"[Sheet: {sheet.title}]")
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_number > 5_000:
                    sections.append("[Sheet truncated at 5000 rows]")
                    break
                sections.append("\t".join("" if value is None else str(value) for value in row))
    finally:
        workbook.close()
    return "\n".join(sections)


def _read_xls(path: Path) -> str:
    workbook = pd.ExcelFile(path, engine="xlrd")
    sections = []
    for sheet_name in workbook.sheet_names:
        frame = pd.read_excel(workbook, sheet_name=sheet_name, nrows=5_000)
        sections.append(f"[Sheet: {sheet_name}]\n{frame.to_csv(index=False)}")
    return "\n\n".join(sections)


READERS = {
    ".txt": _read_text,
    ".md": _read_text,
    ".json": _read_json,
    ".jsonl": _read_jsonl,
    ".csv": _read_csv,
    ".docx": _read_docx,
    ".pdf": _read_pdf,
    ".xlsx": _read_xlsx,
    ".xls": _read_xls,
}

ARCHIVE_EXTENSIONS = {".docx", ".xlsx"}
MAX_ARCHIVE_MEMBERS = 10_000


def _validate_archive(path: Path, settings: Settings) -> None:
    """Reject oversized compressed Office containers before parser expansion."""
    uncompressed_limit = max(
        settings.max_upload_bytes * 4,
        settings.max_extracted_chars * 8,
    )
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ValueError("Office archive contains too many entries")
            if sum(member.file_size for member in members) > uncompressed_limit:
                raise ValueError("Office archive expands beyond the configured safety limit")
    except BadZipFile as exc:
        raise ValueError("invalid Office archive") from exc


def _read_file_sync(file_path: str, settings: Settings) -> dict[str, Any]:
    try:
        path = safe_existing_file(get_session_dir(), file_path)
        extension = path.suffix.lower()
        if extension not in ALLOWED_UPLOAD_EXTENSIONS or extension not in READERS:
            return _file_error("unsupported_type", "unsupported uploaded file type")
        size = path.stat().st_size
        if size > settings.max_upload_bytes:
            return _file_error("too_large", "uploaded file exceeds the configured size limit")
        if extension in ARCHIVE_EXTENSIONS:
            try:
                _validate_archive(path, settings)
            except ValueError as exc:
                return _file_error("unsafe_archive", str(exc))
        extracted = READERS[extension](path)
        content, truncated = _truncate(extracted, settings.max_extracted_chars)
        return {
            "ok": True,
            "name": path.name,
            "extension": extension,
            "size": size,
            "content": content,
            "truncated": truncated,
        }
    except FileNotFoundError:
        return _file_error("not_found", "uploaded file does not exist in this session")
    except SecurityError as exc:
        return _file_error("unsafe_path", str(exc))
    except Exception as exc:
        return _file_error("parse_error", f"file parsing failed: {type(exc).__name__}")


async def read_file_content(file_path: str, *, settings: Settings) -> dict[str, Any]:
    """Read a supported file from the active task session without escaping it."""
    return await asyncio.to_thread(_read_file_sync, file_path, settings)


def make_file_reader_tool(settings: Settings) -> BaseTool:
    async def reader(file_path: str) -> dict[str, Any]:
        """Read one uploaded TXT, MD, JSON, JSONL, CSV, DOCX, PDF, XLSX, or XLS file."""
        return await read_file_content(file_path, settings=settings)

    return StructuredTool.from_function(
        coroutine=reader,
        name="read_file_content",
        description=(
            "Read a supported file already copied into the current task session. Pass only its "
            "session-relative filename; absolute paths and traversal are rejected."
        ),
    )
