import time
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from starlette.websockets import WebSocketDisconnect

from api.server import create_app
from utils.config import Settings


def _settings(tmp_path: Path, **overrides: Any) -> Settings:
    return Settings(
        project_root=tmp_path,
        demo_mode=True,
        max_upload_bytes=overrides.pop("max_upload_bytes", 10 * 1024 * 1024),
        **overrides,
    )


def _receive_until(websocket: Any, terminal_event: str) -> list[dict[str, Any]]:
    events = []
    for _ in range(30):
        payload = websocket.receive_json()
        if payload.get("type") != "monitor_event":
            continue
        events.append(payload)
        if payload.get("event") == terminal_event:
            return events
    raise AssertionError(f"did not receive terminal event: {terminal_event}")


def test_health_and_task_thread_id_contracts(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path))
    with TestClient(app) as client:
        assert client.get("/health").json() == {"status": "ok", "mode": "demo"}

        automatic = client.post("/api/task", json={"query": "Explain agent architecture"})
        assert automatic.status_code == 200
        body = automatic.json()
        assert body["status"] == "started"
        assert len(body["thread_id"]) == 36

        custom = client.post(
            "/api/task",
            json={"query": "Summarize internal policy", "thread_id": "custom-thread"},
        )
        assert custom.status_code == 200
        assert custom.json() == {"status": "started", "thread_id": "custom-thread"}


def test_wildcard_cors_never_enables_credentials(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path, cors_origins=["*"]))
    with TestClient(app) as client:
        response = client.options(
            "/api/task",
            headers={
                "Origin": "https://untrusted.example",
                "Access-Control-Request-Method": "POST",
            },
        )
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "*"
    assert "access-control-allow-credentials" not in response.headers


def test_invalid_task_inputs_are_rejected(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path, max_query_chars=100))
    with TestClient(app) as client:
        assert client.post("/api/task", json={"query": "   "}).status_code == 422
        assert client.post(
            "/api/task", json={"query": "valid", "thread_id": "../../bad"}
        ).status_code == 422
        assert client.post("/api/task", json={"query": "x" * 101}).status_code == 413


def test_upload_multiple_files_sanitizes_names(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path))
    with TestClient(app) as client:
        response = client.post(
            "/api/upload",
            data={"thread_id": "upload-thread"},
            files=[
                ("files", ("safe report.txt", b"hello", "text/plain")),
                ("files", ("data.json", b'{"ok":true}', "application/json")),
            ],
        )
        assert response.status_code == 200
        assert response.json() == {
            "status": "uploaded",
            "files": ["safe_report.txt", "data.json"],
        }
        directory = tmp_path / "updated" / "session_upload-thread"
        assert (directory / "safe_report.txt").read_bytes() == b"hello"
        assert (directory / "data.json").is_file()


def test_upload_rejects_bad_name_extension_and_size_atomically(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path, max_upload_bytes=1_024))
    with TestClient(app) as client:
        traversal = client.post(
            "/api/upload",
            data={"thread_id": "bad-upload"},
            files={"files": ("../../.env", b"secret", "text/plain")},
        )
        assert traversal.status_code == 400

        extension = client.post(
            "/api/upload",
            data={"thread_id": "bad-upload"},
            files={"files": ("payload.exe", b"x", "application/octet-stream")},
        )
        assert extension.status_code == 400

        oversized = client.post(
            "/api/upload",
            data={"thread_id": "bad-upload"},
            files=[
                ("files", ("first.txt", b"first", "text/plain")),
                ("files", ("large.txt", b"x" * 1_025, "text/plain")),
            ],
        )
        assert oversized.status_code == 413
        directory = tmp_path / "updated" / "session_bad-upload"
        assert not (directory / "first.txt").exists()
        assert not list(directory.glob("*.uploading"))


def test_upload_rejects_too_many_files(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path, max_upload_files=2))
    with TestClient(app) as client:
        response = client.post(
            "/api/upload",
            data={"thread_id": "many-files"},
            files=[
                ("files", ("one.txt", b"1", "text/plain")),
                ("files", ("two.txt", b"2", "text/plain")),
                ("files", ("three.txt", b"3", "text/plain")),
            ],
        )
        assert response.status_code == 413


def test_files_and_download_are_output_scoped(tmp_path: Path) -> None:
    output = tmp_path / "output" / "session_files-thread"
    output.mkdir(parents=True)
    (output / "report.md").write_text("# complete report", encoding="utf-8")
    (tmp_path / ".env").write_text("not accessible", encoding="utf-8")

    app = create_app(_settings(tmp_path))
    with TestClient(app) as client:
        listing = client.get("/api/files", params={"path": "session_files-thread"})
        assert listing.status_code == 200
        assert listing.json()[0]["path"] == "session_files-thread/report.md"

        download = client.get(
            "/api/download", params={"path": "session_files-thread/report.md"}
        )
        assert download.status_code == 200
        assert download.content == b"# complete report"

        for unsafe in (
            "../.env",
            "..\\.env",
            "%2e%2e%2f.env",
            "%252e%252e%252f.env",
            "C:\\Windows\\win.ini",
        ):
            assert client.get("/api/files", params={"path": unsafe}).status_code == 403
            assert client.get("/api/download", params={"path": unsafe}).status_code == 403
        assert client.get(
            "/api/download", params={"path": "session_files-thread/missing.md"}
        ).status_code == 404
        assert client.get("/api/files", params={"path": "session_missing"}).status_code == 404


def test_files_and_download_reject_output_symlink_escape(tmp_path: Path) -> None:
    output = tmp_path / "output" / "session_symlink-thread"
    output.mkdir(parents=True)
    outside = tmp_path / "private.txt"
    outside.write_text("private", encoding="utf-8")
    link = output / "report.md"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("symlink creation is unavailable")

    app = create_app(_settings(tmp_path))
    with TestClient(app) as client:
        listing = client.get("/api/files", params={"path": "session_symlink-thread"})
        direct_listing = client.get(
            "/api/files",
            params={"path": "session_symlink-thread/report.md"},
        )
        download = client.get(
            "/api/download",
            params={"path": "session_symlink-thread/report.md"},
        )
    assert listing.status_code == 200 and listing.json() == []
    assert direct_listing.status_code == 403
    assert download.status_code == 403


def test_websocket_ping_and_full_network_markdown_flow(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path))
    thread_id = "network-flow"
    query = "查询人工智能 Agent 最近的发展，并生成一份 Markdown 报告。"
    with TestClient(app) as client, client.websocket_connect(f"/ws/{thread_id}") as websocket:
        websocket.send_text("ping")
        assert websocket.receive_json() == {
            "type": "pong",
            "message": "服务端已收到: ping",
        }
        response = client.post(
            "/api/task", json={"query": query, "thread_id": thread_id}
        )
        assert response.status_code == 200
        events = _receive_until(websocket, "task_result")

        event_names = [event["event"] for event in events]
        assert event_names[0] == "session_created"
        assert "assistant_call" in event_names
        assert "tool_start" in event_names
        assert event_names[-1] == "task_result"
        assert any(
            event["data"].get("assistant_name") == "network_search_agent"
            for event in events
            if event["event"] == "assistant_call"
        )
        assert any(
            event["data"].get("tool_name") == "internet_search"
            for event in events
            if event["event"] == "tool_start"
        )

        session_path = events[0]["data"]["path"]
        files = client.get("/api/files", params={"path": session_path}).json()
        markdown = next(item for item in files if item["name"] == "deep-research-report.md")
        downloaded = client.get("/api/download", params={"path": markdown["path"]})
        assert downloaded.status_code == 200
        assert "# Deep Research Result" in downloaded.text


def test_websocket_disconnect_removes_only_its_connection(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path))
    with TestClient(app) as client:
        with client.websocket_connect("/ws/disconnect-thread") as websocket:
            websocket.send_text("ping")
            websocket.receive_json()
            assert client.portal is not None
            assert client.portal.call(
                app.state.websocket_manager.connection_count, "disconnect-thread"
            ) == 1
        assert client.portal is not None
        assert client.portal.call(
            app.state.websocket_manager.connection_count, "disconnect-thread"
        ) == 0


def test_websocket_rejects_invalid_thread_id(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path))
    with (
        TestClient(app) as client,
        pytest.raises(WebSocketDisconnect) as exc_info,
        client.websocket_connect("/ws/bad%20thread"),
    ):
        pass
    assert exc_info.value.code == 1008


def test_websocket_broadcasts_to_multiple_connections_for_same_thread(
    tmp_path: Path,
) -> None:
    app = create_app(_settings(tmp_path))
    thread_id = "shared-thread"
    with (
        TestClient(app) as client,
        client.websocket_connect(f"/ws/{thread_id}") as socket_one,
        client.websocket_connect(f"/ws/{thread_id}") as socket_two,
    ):
        response = client.post(
            "/api/task",
            json={"query": "Latest agent reliability", "thread_id": thread_id},
        )
        assert response.status_code == 200
        events_one = _receive_until(socket_one, "task_result")
        events_two = _receive_until(socket_two, "task_result")
    assert [event["event"] for event in events_one] == [
        event["event"] for event in events_two
    ]
    assert events_one[-1]["data"]["result"] == events_two[-1]["data"]["result"]


def test_docx_upload_summary_flow(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Quarterly review", level=1)
    document.add_paragraph("Customer retention improved through faster support response.")
    document.save(source)

    app = create_app(_settings(tmp_path))
    thread_id = "document-flow"
    with TestClient(app) as client:
        with source.open("rb") as handle:
            upload = client.post(
                "/api/upload",
                data={"thread_id": thread_id},
                files={
                    "files": (
                        "review.docx",
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert upload.status_code == 200

        with client.websocket_connect(f"/ws/{thread_id}") as websocket:
            assert client.post(
                "/api/task",
                json={"query": "请总结上传的文档", "thread_id": thread_id},
            ).status_code == 200
            events = _receive_until(websocket, "task_result")
        assert any(
            event["data"].get("tool_name") == "read_file_content"
            for event in events
            if event["event"] == "tool_start"
        )
        result = events[-1]["data"]["result"]
        assert "Customer retention improved" in result
        assert (tmp_path / "output" / "session_document-flow" / "review.docx").is_file()


def test_database_agent_flow(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path))
    thread_id = "database-flow"
    with TestClient(app) as client, client.websocket_connect(f"/ws/{thread_id}") as websocket:
        assert client.post(
            "/api/task",
            json={"query": "分析数据库中的产品营收", "thread_id": thread_id},
        ).status_code == 200
        events = _receive_until(websocket, "task_result")
    tool_names = [
        event["data"]["tool_name"]
        for event in events
        if event["event"] == "tool_start"
    ]
    assert tool_names[:3] == ["list_sql_tables", "get_table_data", "execute_sql_query"]
    assert "Delta" in events[-1]["data"]["result"]


def test_parallel_threads_do_not_share_events_or_files(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path))
    with (
        TestClient(app) as client,
        client.websocket_connect("/ws/thread-a") as socket_a,
        client.websocket_connect("/ws/thread-b") as socket_b,
    ):
        client.post(
            "/api/task",
            json={"query": "Latest Agent Alpha report", "thread_id": "thread-a"},
        )
        client.post(
            "/api/task",
            json={"query": "Latest Agent Beta report", "thread_id": "thread-b"},
        )
        events_a = _receive_until(socket_a, "task_result")
        events_b = _receive_until(socket_b, "task_result")

    result_a = events_a[-1]["data"]["result"]
    result_b = events_b[-1]["data"]["result"]
    assert "Agent Alpha" in result_a and "Agent Beta" not in result_a
    assert "Agent Beta" in result_b and "Agent Alpha" not in result_b
    file_a = tmp_path / "output" / "session_thread-a" / "deep-research-report.md"
    file_b = tmp_path / "output" / "session_thread-b" / "deep-research-report.md"
    assert "Agent Alpha" in file_a.read_text(encoding="utf-8")
    assert "Agent Beta" in file_b.read_text(encoding="utf-8")


class FailingRunner:
    async def run(self, query: str, monitor: Any) -> str:
        raise RuntimeError("private provider detail")


def test_runner_failure_emits_safe_error_event(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path), runner=FailingRunner())
    with TestClient(app) as client, client.websocket_connect("/ws/error-thread") as websocket:
        response = client.post(
            "/api/task",
            json={"query": "trigger failure", "thread_id": "error-thread"},
        )
        assert response.status_code == 200
        events = _receive_until(websocket, "error")
    assert events[-1]["data"]["error"] == "Research task failed: RuntimeError"
    assert "private provider detail" not in str(events[-1])

    with TestClient(app) as client:
        for _ in range(20):
            status = client.get("/api/task/error-thread")
            if status.json()["status"] == "failed":
                break
            time.sleep(0.01)
        assert status.json() == {"thread_id": "error-thread", "status": "failed"}


class SlowRunner:
    async def run(self, query: str, monitor: Any) -> str:
        import asyncio

        await asyncio.sleep(0.3)
        return "complete result"


def test_task_endpoint_returns_before_runner_finishes_and_rejects_duplicate(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path), runner=SlowRunner())
    with TestClient(app) as client:
        started = time.monotonic()
        first = client.post(
            "/api/task", json={"query": "slow", "thread_id": "slow-thread"}
        )
        elapsed = time.monotonic() - started
        duplicate = client.post(
            "/api/task", json={"query": "duplicate", "thread_id": "slow-thread"}
        )
        upload = client.post(
            "/api/upload",
            data={"thread_id": "slow-thread"},
            files={"files": ("late.txt", b"late", "text/plain")},
        )
        assert first.status_code == 200
        assert elapsed < 0.2
        assert duplicate.status_code == 409
        assert upload.status_code == 409
