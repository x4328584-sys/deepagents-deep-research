import asyncio
import time
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from pydantic import SecretStr
from reportlab.pdfgen.canvas import Canvas

from api.context import bind_task_context
from tools.markdown_tools import generate_markdown
from tools.mysql_tools import DatabaseService, make_mysql_tools
from tools.pdf_tools import convert_md_to_pdf
from tools.ragflow_tools import RagflowService, make_ragflow_tools
from tools.tavily_tools import internet_search, make_internet_search_tool
from tools.upload_file_read_tool import read_file_content
from utils.config import Settings


class FakeSearchClient:
    async def search(self, query: str, **kwargs: Any) -> dict[str, Any]:
        return {
            "answer": f"answer for {query}",
            "results": [
                {
                    "title": "Primary source",
                    "url": "https://example.test/source",
                    "content": "summary",
                    "raw_content": "full evidence",
                    "score": 0.9,
                    "ignored": "not exposed",
                }
            ],
        }


class FailingSearchClient:
    async def search(self, query: str, **kwargs: Any) -> dict[str, Any]:
        raise RuntimeError("provider exploded")


class SlowSearchClient:
    async def search(self, query: str, **kwargs: Any) -> dict[str, Any]:
        await asyncio.sleep(0.3)
        return {"results": []}


async def test_tavily_demo_and_bounded_counter(tmp_path: Path) -> None:
    settings = Settings(project_root=tmp_path, demo_mode=True, max_search_calls=2)
    with bind_task_context("thread-a", tmp_path):
        first = await internet_search("AI agent architecture", settings=settings)
        second = await internet_search("AI agent evaluation", topic="news", settings=settings)
        third = await internet_search("one too many", settings=settings)
    assert first["ok"] is True and first["mode"] == "demo"
    assert len(first["results"]) >= 3
    assert second["call_number"] == 2
    assert third["error"]["code"] == "search_limit"


async def test_agent_tool_wrappers_are_directly_invokable(tmp_path: Path) -> None:
    settings = Settings(project_root=tmp_path, demo_mode=True)
    mysql_tools = {tool.name: tool for tool in make_mysql_tools(settings)}
    ragflow_tools = {tool.name: tool for tool in make_ragflow_tools(settings)}
    with bind_task_context("thread-tools", tmp_path):
        search = await make_internet_search_tool(settings).ainvoke(
            {"query": "agent reliability", "topic": "general", "max_results": 2}
        )
        tables = await mysql_tools["list_sql_tables"].ainvoke({})
        rows = await mysql_tools["execute_sql_query"].ainvoke(
            {"query": "SELECT name FROM products ORDER BY name"}
        )
        assistants = await ragflow_tools["get_assistant_list"].ainvoke({})
        knowledge = await ragflow_tools["create_ask_delete"].ainvoke(
            {
                "assistant_id": "demo-policy",
                "questions": ["broad", "specific", "exceptions"],
            }
        )
    assert search["ok"] is True and search["results"]
    assert tables["tables"] == ["companies", "products", "sales"]
    assert rows["ok"] is True and rows["rows"]
    assert assistants["ok"] is True and assistants["assistants"]
    assert knowledge["ok"] is True and len(knowledge["answers"]) == 3


async def test_tavily_real_adapter_preserves_source_fields(tmp_path: Path) -> None:
    settings = Settings(
        project_root=tmp_path,
        demo_mode=False,
        tavily_api_key=SecretStr("test-provider-key"),
    )
    with bind_task_context("thread-a", tmp_path):
        result = await internet_search(
            "research topic", settings=settings, client=FakeSearchClient()
        )
    assert result["ok"] is True
    assert result["results"][0]["raw_content"] == "full evidence"
    assert "ignored" not in result["results"][0]


async def test_tavily_provider_failure_is_structured(tmp_path: Path) -> None:
    settings = Settings(
        project_root=tmp_path,
        demo_mode=False,
        tavily_api_key=SecretStr("test-provider-key"),
    )
    with bind_task_context("thread-a", tmp_path):
        result = await internet_search(
            "research topic", settings=settings, client=FailingSearchClient()
        )
    assert result["ok"] is False
    assert result["error"]["code"] == "provider_error"
    assert "provider exploded" not in result["error"]["message"]


async def test_tavily_timeout_is_structured(tmp_path: Path) -> None:
    settings = Settings(
        project_root=tmp_path,
        demo_mode=False,
        tavily_api_key=SecretStr("test-provider-key"),
        external_timeout_seconds=0.01,
    )
    with bind_task_context("thread-a", tmp_path):
        result = await internet_search(
            "research topic", settings=settings, client=SlowSearchClient()
        )
    assert result["error"]["code"] == "timeout"


async def test_demo_database_discovery_preview_and_query(tmp_path: Path) -> None:
    service = DatabaseService(Settings(project_root=tmp_path, demo_mode=True))
    tables = await service.list_tables()
    preview = await service.get_table("products", limit=2)
    query = await service.execute(
        "SELECT name, annual_revenue FROM products ORDER BY annual_revenue DESC LIMIT 2"
    )
    unsafe = await service.execute("DELETE FROM products")
    assert tables["tables"] == ["companies", "products", "sales"]
    assert len(preview["rows"]) == 2 and preview["columns"]
    assert query["rows"][0]["name"] == "Delta"
    assert unsafe["error"]["code"] == "unsafe_sql"


class FakeCursor:
    def __init__(self) -> None:
        self.rows: list[dict[str, Any]] = []

    def __enter__(self) -> "FakeCursor":
        return self

    def __exit__(self, *_: object) -> None:
        return None

    def execute(self, query: str, parameters: object = None) -> None:
        if query == "SHOW TABLES":
            self.rows = [{"Tables_in_company_data": "products"}]
        elif query.startswith("DESCRIBE"):
            self.rows = [{"Field": "id", "Type": "int"}]
        else:
            self.rows = [{"id": 1, "name": "Mock product"}]

    def fetchall(self) -> list[dict[str, Any]]:
        return self.rows

    def fetchmany(self, size: int) -> list[dict[str, Any]]:
        return self.rows[:size]


class FakeConnection:
    def __init__(self) -> None:
        self.closed = False

    def cursor(self) -> FakeCursor:
        return FakeCursor()

    def close(self) -> None:
        self.closed = True


async def test_real_mysql_path_uses_injected_connection(tmp_path: Path) -> None:
    connections: list[FakeConnection] = []

    def factory() -> FakeConnection:
        connection = FakeConnection()
        connections.append(connection)
        return connection

    service = DatabaseService(
        Settings(project_root=tmp_path, demo_mode=False), connection_factory=factory
    )
    tables = await service.list_tables()
    preview = await service.get_table("products")
    rows = await service.execute("SELECT id, name FROM products")
    assert tables["tables"] == ["products"]
    assert preview["columns"][0]["Field"] == "id"
    assert rows["rows"][0]["name"] == "Mock product"
    assert all(connection.closed for connection in connections)


class FakeMessage:
    content = "full knowledge answer"
    reference = [{"document": "policy"}]


class FakeSession:
    id = "session-id"

    def ask(self, question: str, stream: bool = False) -> Any:
        yield FakeMessage()


class FakeChat:
    id = "assistant-id"
    name = "Internal policies"
    description = "Policy corpus"
    prompt_config: dict[str, Any] = {}

    def __init__(self) -> None:
        self.deleted: list[str] = []

    def create_session(self, name: str) -> FakeSession:
        return FakeSession()

    def delete_sessions(self, ids: list[str]) -> None:
        self.deleted.extend(ids)


class FakeRagflow:
    def __init__(self) -> None:
        self.chat = FakeChat()

    def list_chats(self, page: int, page_size: int) -> list[FakeChat]:
        return [self.chat]

    def get_chat(self, assistant_id: str) -> FakeChat:
        assert assistant_id == "assistant-id"
        return self.chat


class SlowRagflow:
    def list_chats(self, page: int, page_size: int) -> list[FakeChat]:
        time.sleep(0.3)
        return []


async def test_ragflow_demo_and_sdk_lifecycle(tmp_path: Path) -> None:
    demo = RagflowService(Settings(project_root=tmp_path, demo_mode=True))
    assistants = await demo.list_assistants()
    answer = await demo.ask_and_cleanup(
        "demo-policy", ["broad policy", "specific control", "exception process"]
    )
    assert assistants["ok"] is True and len(assistants["assistants"]) == 2
    assert len(answer["answers"]) == 3

    fake = FakeRagflow()
    real_settings = Settings(
        project_root=tmp_path,
        demo_mode=False,
        ragflow_api_url="https://ragflow.example.test",
        ragflow_api_key=SecretStr("test-ragflow-key"),
    )
    real = RagflowService(real_settings, sdk_factory=lambda: fake)
    listed = await real.list_assistants()
    queried = await real.ask_and_cleanup(
        "assistant-id", ["broad", "details", "exceptions"]
    )
    assert listed["assistants"][0]["description"] == "Policy corpus"
    assert queried["answers"][0]["content"] == "full knowledge answer"
    assert fake.chat.deleted == ["session-id"]


async def test_ragflow_missing_configuration_degrades(tmp_path: Path) -> None:
    service = RagflowService(Settings(project_root=tmp_path, demo_mode=False))
    result = await service.list_assistants()
    assert result["ok"] is False
    assert result["error"]["code"] == "not_configured"


async def test_ragflow_timeout_degrades(tmp_path: Path) -> None:
    settings = Settings(
        project_root=tmp_path,
        demo_mode=False,
        ragflow_api_url="https://ragflow.example.test",
        ragflow_api_key=SecretStr("test-ragflow-key"),
        external_timeout_seconds=0.01,
    )
    service = RagflowService(settings, sdk_factory=SlowRagflow)
    result = await service.list_assistants()
    assert result["error"]["code"] == "timeout"


def _create_supported_files(root: Path) -> list[str]:
    (root / "notes.txt").write_text("hello text", encoding="utf-8")
    (root / "notes.md").write_text("# Hello\nMarkdown", encoding="utf-8")
    (root / "data.json").write_text('{"name":"demo"}', encoding="utf-8")
    (root / "data.jsonl").write_text('{"id":1}\n{"id":2}\n', encoding="utf-8")
    (root / "data.csv").write_text("name,value\nalpha,1\n", encoding="utf-8")

    document = Document()
    document.add_heading("Uploaded document", level=1)
    document.add_paragraph("Document body")
    document.save(root / "document.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "value"])
    sheet.append(["alpha", 1])
    workbook.save(root / "workbook.xlsx")

    canvas = Canvas(str(root / "document.pdf"))
    canvas.drawString(72, 720, "PDF document body")
    canvas.save()
    return [
        "notes.txt",
        "notes.md",
        "data.json",
        "data.jsonl",
        "data.csv",
        "document.docx",
        "workbook.xlsx",
        "document.pdf",
    ]


async def test_file_reader_supports_document_formats(tmp_path: Path) -> None:
    settings = Settings(project_root=tmp_path)
    filenames = _create_supported_files(tmp_path)
    with bind_task_context("thread-a", tmp_path):
        results = await asyncio.gather(
            *(read_file_content(name, settings=settings) for name in filenames)
        )
    assert all(result["ok"] is True for result in results)
    assert "Document body" in results[5]["content"]
    assert "alpha" in results[6]["content"]
    assert "PDF document body" in results[7]["content"]


async def test_file_reader_blocks_escape_and_oversize(tmp_path: Path) -> None:
    (tmp_path / "large.txt").write_bytes(b"x" * 1_025)
    settings = Settings(project_root=tmp_path, max_upload_bytes=1_024)
    with bind_task_context("thread-a", tmp_path):
        traversal = await read_file_content("../secret.txt", settings=settings)
        large = await read_file_content("large.txt", settings=settings)
    assert traversal["error"]["code"] == "unsafe_path"
    assert large["error"]["code"] == "too_large"


async def test_file_reader_rejects_office_archive_expansion_bomb(tmp_path: Path) -> None:
    archive_path = tmp_path / "bomb.docx"
    with ZipFile(archive_path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * (5 * 1024 * 1024))
    settings = Settings(
        project_root=tmp_path,
        max_upload_bytes=1024 * 1024,
        max_extracted_chars=1_000,
    )
    assert archive_path.stat().st_size < settings.max_upload_bytes
    with bind_task_context("thread-a", tmp_path):
        result = await read_file_content("bomb.docx", settings=settings)
    assert result["error"]["code"] == "unsafe_archive"


async def test_markdown_then_pdf_generation(tmp_path: Path) -> None:
    with bind_task_context("thread-a", tmp_path):
        markdown = await generate_markdown(
            "研究 report.md",
            "# 研究报告\n\n这是一个完整的中文测试报告，包含经过验证的信息和结论。",
        )
        pdf = await convert_md_to_pdf(markdown["path"])
    assert markdown["ok"] is True
    assert (tmp_path / markdown["path"]).read_text(encoding="utf-8").startswith("# 研究报告")
    assert pdf["ok"] is True
    assert (tmp_path / pdf["path"]).read_bytes().startswith(b"%PDF")


async def test_placeholder_markdown_is_rejected(tmp_path: Path) -> None:
    with bind_task_context("thread-a", tmp_path):
        result = await generate_markdown("report.md", "TODO")
    assert result["ok"] is False
    assert not (tmp_path / "report.md").exists()
